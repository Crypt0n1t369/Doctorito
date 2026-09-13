from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


OUT = Path("/Users/kristaps/Documents/New project/output/atbalsts-kiberdrosibas-stavoklis-2026-09-10.docx")

BLACK = "000000"
NAVY = "183B56"
BLUE = "2F6B8A"
PALE_BLUE = "EEF5F8"
PALE_GRAY = "F5F6F7"
MID_GRAY = "6B7280"
LIGHT_GRAY = "D9D9D9"
GREEN = "1F6F5C"
AMBER = "A56A16"
RED = "A13D3D"
WHITE = "FFFFFF"
FONT = "Liberation Sans"
MONO = "Liberation Mono"


def set_run_font(run, name=FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_mm):
    cell.width = Mm(width_mm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_mm * 56.6929)))
    tc_w.set(qn("w:type"), "dxa")


def style_paragraph(paragraph, before=0, after=3, line=1.02, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next


def remove_paragraph_borders(element):
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_hyperlink(paragraph, text, url, size=8.2):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    font_el = OxmlElement("w:rFonts")
    font_el.set(qn("w:ascii"), FONT)
    font_el.set(qn("w:hAnsi"), FONT)
    r_pr.extend([font_el, color, underline, size_el])
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run("Lapa ")
    set_run_font(run, size=7.8, color=MID_GRAY)
    for field_name in ("PAGE", "NUMPAGES"):
        if field_name == "NUMPAGES":
            separator = paragraph.add_run(" no ")
            set_run_font(separator, size=7.8, color=MID_GRAY)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), field_name)
        run_el = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), MID_GRAY)
        size_el = OxmlElement("w:sz")
        size_el.set(qn("w:val"), "16")
        r_pr.extend([color, size_el])
        run_el.append(r_pr)
        text_el = OxmlElement("w:t")
        text_el.text = "1" if field_name == "PAGE" else "2"
        run_el.append(text_el)
        field.append(run_el)
        paragraph._p.append(field)


def score_color(score):
    if score >= 80:
        return GREEN
    if score >= 65:
        return BLUE
    if score >= 50:
        return AMBER
    return RED


def add_progress(paragraph, score, compact=False):
    filled = max(0, min(10, round(score / 10)))
    size = 8.5 if compact else 10.5
    run = paragraph.add_run("■" * filled)
    set_run_font(run, name=MONO, size=size, bold=True, color=score_color(score))
    run = paragraph.add_run("■" * (10 - filled))
    set_run_font(run, name=MONO, size=size, bold=True, color="DDE2E7")
    run = paragraph.add_run(f"  {score}%")
    set_run_font(run, size=8.5 if compact else 10.5, bold=True, color=BLACK)


def add_heading(doc, text, level=1, before=7, after=3):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(run, size=13 if level == 1 else 10.5, bold=True, color=BLACK)
    style_paragraph(paragraph, before=before, after=after, line=1.0, keep_next=True)
    return paragraph


def add_bullet(doc, lead, text, size=9.2, after=1.6):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Mm(4.5)
    paragraph.paragraph_format.first_line_indent = Mm(-2.5)
    style_paragraph(paragraph, after=after, line=1.02)
    run = paragraph.add_run(lead)
    set_run_font(run, size=size, bold=True, color=BLACK)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=BLACK)
    return paragraph


def add_module_table(doc, module_rows):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = ("Aspekts", "Gatavība", "Pierādīts tagad", "Atlikušais darbs")
    widths = (34, 36, 52, 60)
    for idx, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=65, bottom=65, start=75, end=75)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(p, after=0, line=1.0)
        r = p.add_run(text)
        set_run_font(r, size=8.1, bold=True, color=WHITE)
    repeat_header(table.rows[0])

    for idx, (aspect, score, current_state, remaining) in enumerate(module_rows):
        row = table.add_row()
        prevent_row_split(row)
        if idx % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for c_idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[c_idx])
            set_cell_margins(cell, top=48, bottom=48, start=62, end=62)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = row.cells[0].paragraphs[0]
        style_paragraph(p, after=0, line=1.0)
        r = p.add_run(aspect)
        set_run_font(r, size=7.7, bold=True, color=BLACK)
        p = row.cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        add_progress(p, score, compact=True)
        for c_idx, text in ((2, current_state), (3, remaining)):
            p = row.cells[c_idx].paragraphs[0]
            style_paragraph(p, after=0, line=1.0)
            r = p.add_run(text)
            set_run_font(r, size=7.55, color=BLACK)

    return table


doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(10)
section.bottom_margin = Mm(10)
section.left_margin = Mm(13)
section.right_margin = Mm(13)
section.header_distance = Mm(4)
section.footer_distance = Mm(4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.03

for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.color.rgb = RGBColor.from_string(BLACK)
    remove_paragraph_borders(style._element)

title = doc.add_paragraph(style="Title")
remove_paragraph_borders(title._p)
title_run = title.add_run("Atbalsts platformas kiberdrošības stāvoklis")
set_run_font(title_run, size=20, bold=True, color=BLACK)
style_paragraph(title, after=1, line=1.0, keep_next=True)

subtitle = doc.add_paragraph(style="Subtitle")
subtitle_run = subtitle.add_run("Stāvoklis 2026. gada 10. septembrī")
set_run_font(subtitle_run, size=9.4, color=MID_GRAY)
style_paragraph(subtitle, after=5, line=1.0, keep_next=True)

intro = doc.add_paragraph()
style_paragraph(intro, after=4, line=1.08)
run = intro.add_run("Galvenais secinājums  ")
set_run_font(run, size=10.2, bold=True, color=BLACK)
run = intro.add_run(
    "Platforma ir gatava skaidri marķētai demonstrācijai, un papildu tehnisku prezentācijas šķēršļu nav. PR #73 ir apvienots pēc sekmīgas CI pārbaudes, "
    "publiskajā vidē darbojas tieši apstiprinātais galvenā zara laidiens 664b5e5, sensitīvā diagnostika ir slēgta, bet atgriešanās un abu procesu automātiska atjaunošanās ir praktiski pārbaudīta. "
    "Tas vēl nav pilnīgas gatavības darbam ražošanas vidē vai juridiskas atbilstības apliecinājums."
)
set_run_font(run, size=10.2, color=BLACK)

score_table = doc.add_table(rows=2, cols=2)
score_table.alignment = WD_TABLE_ALIGNMENT.CENTER
score_table.autofit = False
set_table_borders(score_table, color="FFFFFF", size="0")
for row in score_table.rows:
    set_cell_width(row.cells[0], 91)
    set_cell_width(row.cells[1], 91)
    for cell in row.cells:
        set_cell_margins(cell, top=35, bottom=35, start=40, end=40)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

for idx, (label, score) in enumerate((("Prezentācijas gatavība", 93), ("Ražošanas vides drošības gatavība", 66))):
    cell = score_table.rows[0].cells[idx]
    p = cell.paragraphs[0]
    style_paragraph(p, after=0, line=1.0)
    r = p.add_run(label)
    set_run_font(r, size=9.2, bold=True, color=BLACK)
    p = score_table.rows[1].cells[idx].paragraphs[0]
    style_paragraph(p, after=0, line=1.0)
    add_progress(p, score)

note = doc.add_paragraph()
style_paragraph(note, after=4, line=1.0)
run = note.add_run(
    "Procenti ir kontroles gatavības novērtējums pēc koda, publiskās vides, procesa un pierādījumu stāvokļa. Tie nav sertifikācija, incidenta varbūtība vai juridisks atzinums."
)
set_run_font(run, size=7.8, italic=True, color=MID_GRAY)

add_heading(doc, "Pārbaudīts pirms prezentācijas", level=1, before=3, after=2)

verified_rows = [
    ("Izlaidums", "PR #73 apvienots pēc sekmīgas CI pārbaudes; publiskajā vidē izvērsts galvenā zara SHA 664b5e58bde2fe39287b806c5606eb573a78c1ae."),
    ("Publiskā virsma", "Galvenās lapas atbild ar 200; /api/health izdod tikai ok un būvējuma SHA ar no-store; publiskajā notikumu API nav iekšējo lauku."),
    ("Piekļuves kontrole", "Anonīms /api/center/health un POST /api/live-data/refresh saņem 401."),
    ("Atkopšana", "Izmēģināta pāreja uz iepriekšējo SHA 6636f6f un atjaunošana uz 664b5e5; abas versijas publiski identificējās pareizi."),
    ("Pieejamība", "Lietotne un Cloudflare Tunnel darbojas ieslēgtos systemd lietotāja servisos ar Restart=always un pēc piespiedu procesa kļūmes atjaunojās automātiski."),
]

verified = doc.add_table(rows=1, cols=2)
verified.alignment = WD_TABLE_ALIGNMENT.CENTER
verified.autofit = False
set_table_borders(verified)
for i, text in enumerate(("Kontrole", "Pierādītais rezultāts")):
    cell = verified.rows[0].cells[i]
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=70, bottom=70, start=90, end=90)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    style_paragraph(p, after=0, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=8.5, bold=True, color=WHITE)
set_cell_width(verified.rows[0].cells[0], 35)
set_cell_width(verified.rows[0].cells[1], 147)
repeat_header(verified.rows[0])

for idx, (control, result) in enumerate(verified_rows):
    row = verified.add_row()
    prevent_row_split(row)
    if idx % 2 == 1:
        for cell in row.cells:
            set_cell_shading(cell, PALE_BLUE)
    set_cell_width(row.cells[0], 35)
    set_cell_width(row.cells[1], 147)
    for cell in row.cells:
        set_cell_margins(cell, top=58, bottom=58, start=85, end=85)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = row.cells[0].paragraphs[0]
    style_paragraph(p, after=0, line=1.0)
    r = p.add_run(control)
    set_run_font(r, size=8.3, bold=True, color=BLACK)
    p = row.cells[1].paragraphs[0]
    style_paragraph(p, after=0, line=1.02)
    r = p.add_run(result)
    set_run_font(r, size=8.3, color=BLACK)

add_heading(doc, "Pašreizējā arhitektūra", level=1, before=5, after=2)
arch = doc.add_paragraph()
style_paragraph(arch, after=2, line=1.04)
r = arch.add_run("Internets → Cloudflare starpniekserveris → Cloudflare Tunnel → lokāls izcelsmes servers → Vite priekšskatījuma serveris")
set_run_font(r, size=9.2, bold=True, color=NAVY)

arch2 = doc.add_paragraph()
style_paragraph(arch2, after=2, line=1.04)
r = arch2.add_run(
    "Tunelis samazina tiešu servera ekspozīciju, bet neaizstāj WAF, hosta aizsardzību vai vairāku instanču pieejamību. Cloudflare Worker pastāv, taču nav publiskā domēna izcelsmes serveris."
)
set_run_font(r, size=9.0, color=BLACK)

add_heading(doc, "Prezentācijā norādāmie ierobežojumi", level=1, before=4, after=1)
add_bullet(doc, "Apzināti pieļaujams  ", "kopīgā Centra parole, sintētiski dati un Vite priekšskatījuma serveris tikai marķētā demonstrācijas režīmā.", size=8.9, after=1)
add_bullet(doc, "Nedrīkst apgalvot  ", "pilnu produkcijas gatavību, aktīvus pielāgotus Cloudflare WAF noteikumus, darbību piesaisti individuāliem operatoriem vai ISO/IEC 27001 sertifikāciju.", size=8.9, after=0)

add_heading(doc, "Izolēto aspektu drošības stāvoklis", level=1, before=0, after=3)

modules = [
    ("Publiskā virsma", 92, "Minimāla /api/health atbilde, noteikts publisko lauku saraksts un sekmīgas anonīmās pārbaudes.", "Perimetra pieprasījumu ierobežojumi un pārējo publisko datu projekciju periodiska pārbaude."),
    ("Izlaidumi un atkopšana", 90, "Precīzs SHA, nemaināmi laidieni, restartēšana un atgriešanās pārbaudīta.", "Aizstāt Vite priekšskatījuma serveri; ieviest otru izcelsmes instanci vai pārbaudītu Worker vidi."),
    ("MI un ārkārtas loģika", 82, "Ārkārtas atbilde pirms MI; publicē tikai cilvēks; shēmas un avoti.", "Promptu injekcijas un pretinieka scenāriju pārbaudes, centralizēta telemetrija un drošs pakalpojuma atslēgšanas režīms."),
    ("Repozitorijs un piegādes ķēde", 74, "CI, 32 pārbaudes, būvējums, kritisku atkarību vārti, Dependabot un CODEOWNERS.", "Galvenā zara aizsardzība, drošības brīdinājumi un augsto un vidējo atkarību sasniedzamības analīze."),
    ("Dati un Supabase", 65, "Servera atslēgas, noturīga glabātuve, lauku minimizācija un validācija.", "Datu karte, ietekmes uz datu aizsardzību novērtējums, glabāšana un dzēšana, minimālas datubāzes lomas, audits un atjaunošanas tests."),
    ("Identitāte un sesijas", 58, "Droši sīkfaili, izcelsmes pārbaudes, iekšējā diagnostika pieejama tikai autentificētā sesijā.", "Individuāli konti, MFA, sesiju atsaukšana, ierīču pārvaldība un darbību attiecināmība."),
    ("Faili", 55, "Privāta glabāšana, īpašnieka, izmēra un deklarētā tipa pārbaudes.", "Karantīna, ļaunprogrammatūras skenēšana un droša satura pārkodēšana."),
    ("Cloudflare perimetrs", 45, "Cloudflare reversais starpniekserveris un tunelis; lietotnes pieprasījumu ierobežojumi; anonīma datu atsvaidzināšana ir slēgta.", "R1 MI API un R3 pieteikšanās ātruma ierobežojumi, WAF un Security Events pārbaude."),
    ("Misijas un brīvprātīgie", 45, "Cilvēka apstiprinājums publicēšanai, versiju kontrole, atkārtotu pieprasījumu drošums un pamata auditācijas pieraksti.", "Kompetences un riska pārbaudes, četru acu princips, piekrišana un tās atsaukšana."),
    ("Uzraudzība un incidenti", 55, "Minimāla veselības pārbaude, uzraudzīti procesi un pārbaudīta atgriešanās.", "Pieejamības un ļaunprātīgas izmantošanas brīdinājumi, incidentu plāns, rezerves kopijas atjaunošana un galda mācības."),
]

add_module_table(doc, modules[:5])

second_section = doc.add_section(WD_SECTION.NEW_PAGE)
second_section.page_width = Mm(210)
second_section.page_height = Mm(297)
second_section.top_margin = Mm(10)
second_section.bottom_margin = Mm(10)
second_section.left_margin = Mm(13)
second_section.right_margin = Mm(13)
second_section.header_distance = Mm(4)
second_section.footer_distance = Mm(4)

add_heading(doc, "Izolēto aspektu drošības stāvoklis — turpinājums", level=2, before=0, after=3)
add_module_table(doc, modules[5:])

add_heading(doc, "Svarīgākie atlikušā darba īpašnieki", level=1, before=4, after=1)
add_bullet(doc, "Cloudflare zonas administrators — ieteicamais termiņš 2026. gada 17. septembris  ", "ieviest R1 MI API un R3 pieteikšanās ātruma ierobežojumus ar Managed Challenge, pēc tam septiņas dienas pārskatīt Security Events žurnālu. Ja prezentācija ir agrāk, ierobežojums jānorāda un lietotnes aizsardzība jāsaglabā kā kompensējoša kontrole.", size=8.1, after=0.7)
add_bullet(doc, "GitHub repozitorija administrators — ieteicamais termiņš 2026. gada 17. septembris  ", "aizsargāt galveno zaru main ar obligātu PR, CI pārbaudi verify un CODEOWNERS apstiprinājumu; aizliegt piespiedu pārrakstīšanu un dzēšanu; ieslēgt pieejamos Dependabot un slepeno datu noplūdes brīdinājumus.", size=8.1, after=0.7)
add_bullet(doc, "Infrastruktūras īpašnieks — ieteicamais termiņš 2026. gada 10. oktobris  ", "aizstāt Vite priekšskatījuma serveri ar atbalstītu ražošanas izpildvidi vai pārbaudīt Worker migrāciju pirmsražošanas domēnā; pievienot pieejamības brīdinājumus.", size=8.1, after=0.7)
add_bullet(doc, "Produkta privātuma un drošības īpašnieki — pirms personas datu pilota  ", "pabeigt datu karti un novērtējumu par ietekmi uz datu aizsardzību (NIDA), ieviest individuālus lietotāju kontus un MFA, sesiju atsaukšanu, failu karantīnu, misiju riska pārbaudes un rezerves kopiju atjaunošanas mācības.", size=8.1, after=1)

attack = doc.add_paragraph()
style_paragraph(attack, before=2, after=1, line=1.02)
r = attack.add_run("Galvenie uzbrukuma ceļi  ")
set_run_font(r, size=8.4, bold=True, color=BLACK)
r = attack.add_run("Programmatūrā tie ir API datu izguve, pieprasījumu plūdi un izmaksu mākslīga palielināšana, sesiju zādzība, promptu injekcija, ļaunprātīgi faili un piegādes ķēdes risks. Cilvēku un procesu pusē tie ir pikšķerēšana, kopīgi piekļuves dati, uzdošanās par partneri, iekšēja kļūda, nogurums un neizmēģināta atkopšana pēc incidenta.")
set_run_font(r, size=8.4, color=BLACK)

standards = doc.add_paragraph()
style_paragraph(standards, after=0, line=1.02)
r = standards.add_run("Standartu ceļš  ")
set_run_font(r, size=8.4, bold=True, color=BLACK)
r = standards.add_run("Personas datu apstrādei pirmais obligātais slānis ir VDAR. Atsevišķi jānosaka Nacionālā kiberdrošības likuma, NIS2 un MK noteikumu Nr. 397 piemērojamība, Kibernoturības akta piemērojamība katram tirgū laistam produktam ar digitāliem elementiem un MI akta prasības pēc MI faktiskās lomas. Ieviešanai izmantot NIST CSF 2.0 un OWASP ASVS 5.0; ISO/IEC 27001 sertifikāciju apsvērt tad, kad to prasa tirgus vai iepirkums.")
set_run_font(r, size=8.4, color=BLACK)

sources = doc.add_paragraph()
style_paragraph(sources, after=0, line=1.0)
r = sources.add_run("Primārie avoti: ")
set_run_font(r, size=7.2, bold=True, color=MID_GRAY)
source_links = (
    ("VDAR", "https://eur-lex.europa.eu/eli/reg/2016/679/oj"),
    ("Nacionālās kiberdrošības likums", "https://likumi.lv/ta/id/353390-nacionalas-kiberdrosibas-likums"),
    ("MK Nr. 397", "https://likumi.lv/ta/id/361481-minimalas-kiberdrosibas-prasibas"),
    ("Kibernoturības akts", "https://digital-strategy.ec.europa.eu/en/policies/cra-summary"),
    ("MI akts", "https://eur-lex.europa.eu/eli/reg/2024/1689/oj"),
    ("NIST CSF 2.0", "https://www.nist.gov/cyberframework"),
    ("OWASP ASVS 5.0", "https://owasp.org/www-project-application-security-verification-standard/"),
)
for idx, (label, url) in enumerate(source_links):
    if idx:
        r = sources.add_run(" · ")
        set_run_font(r, size=7.2, color=MID_GRAY)
    add_hyperlink(sources, label, url, size=7.2)

add_heading(doc, "Produktu palaišanas lēmumi", level=2, before=3, after=1)
add_bullet(doc, "Publiskā informācija un tikai lasāma karte  ", "drošākais pirmais atsevišķais produkts; pirms palaišanas publiskā vidē ieviest perimetra pieprasījumu ierobežojumus un atbalstītu izpildvidi.", size=8.0, after=0.5)
add_bullet(doc, "Iedzīvotāju ziņojumi  ", "palaist tikai pēc datu kartes, glabāšanas un dzēšanas noteikumiem, centralizētiem pieprasījumu ierobežojumiem un failu karantīnas.", size=8.0, after=0.5)
add_bullet(doc, "Vadības centrs  ", "saglabāt demonstrācijas režīmā līdz individuāliem kontiem, MFA, sesiju atsaukšanai un pret viltojumiem aizsargātai audita uzskaitei.", size=8.0, after=0.5)
add_bullet(doc, "Misijas un brīvprātīgo nosūtīšana  ", "izolēt vai neizlaist līdz kompetences un riska pārbaudēm, piekrišanai un četru acu principam.", size=8.0, after=0.5)
add_bullet(doc, "MI starpslānis  ", "saglabāt tikai padomdevēja lomā ar minimizētu kontekstu un bez publicēšanas vai cilvēku nosūtīšanas tiesībām.", size=8.0, after=0)

doc.core_properties.title = "Atbalsts platformas kiberdrošības stāvoklis"
doc.core_properties.subject = "Divu A4 lapu prezentācijas drošības kopsavilkums"
doc.core_properties.author = "Sortium"
doc.core_properties.keywords = "Atbalsts kiberdrošība Cloudflare VDAR NIS2 OWASP NIST"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
