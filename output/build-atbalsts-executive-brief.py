from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


OUT = Path("/Users/kristaps/Documents/New project/output/atbalsts-kiberdrosibas-vadibas-kopsavilkums-2026-09-10.docx")

BLACK = "000000"
NAVY = "173A53"
BLUE = "2F6B8A"
GREEN = "1F6F5C"
AMBER = "A56A16"
RED = "A13D3D"
GRAY = "66717C"
LIGHT = "DDE3E8"
PALE = "F3F6F8"
WHITE = "FFFFFF"
FONT = "Liberation Sans"
MONO = "Liberation Mono"


def set_run(run, size=10.5, bold=False, color=BLACK, italic=False, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_p(p, before=0, after=4, line=1.08, keep=False):
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line
    f.keep_with_next = keep


def remove_borders(element):
    ppr = element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def set_cell_width(cell, mm):
    cell.width = Mm(mm)
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(mm * 56.6929)))
    tcw.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=50, start=70, bottom=50, end=70):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def hide_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def shade_cell(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), color)


def bottom_rule(p, color=LIGHT, size="8"):
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def heading(doc, text, size=14, before=8, after=4):
    p = doc.add_paragraph(style="Heading 1")
    remove_borders(p._p)
    style_p(p, before=before, after=after, line=1.0, keep=True)
    r = p.add_run(text)
    set_run(r, size=size, bold=True, color=BLACK)
    return p


def add_hyperlink(p, text, url, size=7.4):
    rel = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.extend([fonts, color, underline, sz])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    p._p.append(hyperlink)


def score_color(score):
    if score >= 80:
        return GREEN
    if score >= 65:
        return BLUE
    if score >= 50:
        return AMBER
    return RED


def progress(p, score, size=9.5):
    filled = max(0, min(10, round(score / 10)))
    r = p.add_run("━" * filled)
    set_run(r, size=size, bold=True, color=score_color(score), name=MONO)
    r = p.add_run("━" * (10 - filled))
    set_run(r, size=size, bold=True, color=LIGHT, name=MONO)


def status_metric(cell, label, score, explanation):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell, top=70, start=40, bottom=70, end=130)
    p = cell.paragraphs[0]
    style_p(p, after=1, line=1.0)
    r = p.add_run(f"{score}%")
    set_run(r, size=22, bold=True, color=score_color(score))
    r = p.add_run(f"  {label}")
    set_run(r, size=10.5, bold=True)
    p = cell.add_paragraph()
    style_p(p, after=1, line=1.0)
    progress(p, score, size=10)
    p = cell.add_paragraph()
    style_p(p, after=0, line=1.04)
    r = p.add_run(explanation)
    set_run(r, size=8.5, color=GRAY)


def check_item(cell, title, text):
    p = cell.add_paragraph()
    style_p(p, after=1, line=1.04)
    r = p.add_run("✓  ")
    set_run(r, size=11, bold=True, color=GREEN)
    r = p.add_run(title)
    set_run(r, size=10.2, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Mm(5)
    style_p(p, after=5, line=1.08)
    r = p.add_run(text)
    set_run(r, size=9.2, color=GRAY)


def phase(cell, number, title, status, color, text):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(cell, top=70, start=80, bottom=50, end=120)
    p = cell.paragraphs[0]
    style_p(p, after=2, line=1.0)
    r = p.add_run(f"{number:02d}  ")
    set_run(r, size=9, bold=True, color=GRAY)
    r = p.add_run(title)
    set_run(r, size=10.3, bold=True)
    p = cell.add_paragraph()
    style_p(p, after=3, line=1.0)
    r = p.add_run(status.upper())
    set_run(r, size=8.2, bold=True, color=color)
    p = cell.add_paragraph()
    style_p(p, after=0, line=1.06)
    r = p.add_run(text)
    set_run(r, size=8.9, color=BLACK)


def aspect(cell, label, score, note):
    p = cell.add_paragraph()
    style_p(p, before=0, after=0.5, line=1.0)
    r = p.add_run(label)
    set_run(r, size=9.3, bold=True)
    r = p.add_run(f"  {score}%")
    set_run(r, size=9.3, bold=True, color=score_color(score))
    p = cell.add_paragraph()
    style_p(p, after=0.5, line=1.0)
    progress(p, score, size=7.7)
    p = cell.add_paragraph()
    style_p(p, after=5.2, line=1.04)
    r = p.add_run(note)
    set_run(r, size=8.15, color=GRAY)


def practice_item(cell, symbol, title, text, color):
    p = cell.add_paragraph()
    style_p(p, after=1, line=1.02)
    r = p.add_run(f"{symbol}  ")
    set_run(r, size=9.4, bold=True, color=color)
    r = p.add_run(title)
    set_run(r, size=9.4, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Mm(4.8)
    style_p(p, after=5.5, line=1.07)
    r = p.add_run(text)
    set_run(r, size=8.5, color=GRAY)


def compact_aspect(cell, label, score):
    p = cell.add_paragraph()
    style_p(p, after=0.5, line=1.0)
    r = p.add_run(label)
    set_run(r, size=8.8, bold=True)
    r = p.add_run(f"  {score}%  ")
    set_run(r, size=8.8, bold=True, color=score_color(score))
    progress(p, score, size=6.8)
    style_p(p, after=4, line=1.0)


def action(cell, number, title, owner, deadline, text, color):
    p = cell.add_paragraph()
    style_p(p, after=1, line=1.0)
    r = p.add_run(f"{number:02d}")
    set_run(r, size=15, bold=True, color=color)
    r = p.add_run(f"  {title}")
    set_run(r, size=9.7, bold=True)
    p = cell.add_paragraph()
    style_p(p, after=2, line=1.0)
    r = p.add_run(f"{owner}  ·  {deadline}")
    set_run(r, size=7.8, bold=True, color=GRAY)
    p = cell.add_paragraph()
    style_p(p, after=5, line=1.07)
    r = p.add_run(text)
    set_run(r, size=8.5)


def attack_paths(doc):
    risk = doc.add_table(rows=1, cols=2)
    risk.alignment = WD_TABLE_ALIGNMENT.CENTER
    risk.autofit = False
    hide_table_borders(risk)
    for c in risk.rows[0].cells:
        set_cell_width(c, 88)
        set_cell_margins(c, top=20, start=0, bottom=0, end=140)

    p = risk.rows[0].cells[0].paragraphs[0]
    style_p(p, after=1, line=1.0)
    r = p.add_run("PROGRAMMATŪRA")
    set_run(r, size=8, bold=True, color=RED)
    p = risk.rows[0].cells[0].add_paragraph()
    style_p(p, after=0, line=1.08)
    r = p.add_run("API datu izguve, pieprasījumu plūdi un izmaksu palielināšana, sesiju zādzība, promptu injekcija, ļaunprātīgi faili un piegādes ķēdes kompromitēšana.")
    set_run(r, size=8.6)

    p = risk.rows[0].cells[1].paragraphs[0]
    style_p(p, after=1, line=1.0)
    r = p.add_run("CILVĒKI UN PROCESS")
    set_run(r, size=8, bold=True, color=RED)
    p = risk.rows[0].cells[1].add_paragraph()
    style_p(p, after=0, line=1.08)
    r = p.add_run("Pikšķerēšana, kopīgi piekļuves dati, uzdošanās par partneri, kļūdaina publicēšana, iekšēja ļaunprātīga rīcība, nogurums un nepārbaudīta atkopšanas kārtība.")
    set_run(r, size=8.6)


doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(13)
section.bottom_margin = Mm(13)
section.left_margin = Mm(15)
section.right_margin = Mm(15)
section.header_distance = Mm(5)
section.footer_distance = Mm(6)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(10.3)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.08

for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
    s = styles[name]
    s.font.name = FONT
    s._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    s._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    s.font.color.rgb = RGBColor.from_string(BLACK)
    remove_borders(s._element)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_p(footer, after=0, line=1.0)
r = footer.add_run("Atbalsts kiberdrošības vadības kopsavilkums  ·  2026. gada 10. septembris")
set_run(r, size=7.3, color=GRAY)

# PAGE 1
title = doc.add_paragraph(style="Title")
remove_borders(title._p)
style_p(title, after=1, line=1.0, keep=True)
r = title.add_run("Atbalsts platformas drošības gatavība")
set_run(r, size=23, bold=True)

subtitle = doc.add_paragraph(style="Subtitle")
style_p(subtitle, after=9, line=1.0, keep=True)
r = subtitle.add_run("Vadības kopsavilkums krīzes centra vadībai")
set_run(r, size=10.5, color=GRAY)

lead = doc.add_paragraph()
style_p(lead, after=8, line=1.14)
r = lead.add_run("Secinājums  ")
set_run(r, size=11.4, bold=True)
r = lead.add_run(
    "Platformu var droši demonstrēt, ja tā ir skaidri marķēta kā demonstrācijas vide. "
    "Apstiprinātais laidiens, minimāla publiskā diagnostika, piekļuves kontrole, procesu atjaunošanās un atgriešanās uz iepriekšējo versiju ir pārbaudīta. "
    "Reālai krīzes vadībai vēl jāievieš individuāli lietotāju konti, perimetra aizsardzība, darbības nepārtrauktība un obligātas cilvēku drošības pārbaudes."
)
set_run(r, size=11.1)

metrics = doc.add_table(rows=1, cols=2)
metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
metrics.autofit = False
hide_table_borders(metrics)
for c in metrics.rows[0].cells:
    set_cell_width(c, 88)
status_metric(metrics.rows[0].cells[0], "Demonstrācijas gatavība", 93, "Augsta. Būtiskākie demonstrācijas kontroles punkti ir praktiski pārbaudīti; ārējie administratoru darbi ir atklāti norādīti.")
status_metric(metrics.rows[0].cells[1], "Ekspluatācijas drošības gatavība", 66, "Vidēja. Atlikušie 34 procentpunkti un to ieviešanas plāns izskaidrots otrajā lapā; tas nav audits vai sertifikācija.")

h = heading(doc, "Ko var droši parādīt šodien", size=14, before=8, after=2)
bottom_rule(h, color=LIGHT, size="6")
checks = doc.add_table(rows=1, cols=2)
checks.alignment = WD_TABLE_ALIGNMENT.CENTER
checks.autofit = False
hide_table_borders(checks)
for c in checks.rows[0].cells:
    set_cell_width(c, 88)
    set_cell_margins(c, top=20, start=0, bottom=0, end=90)
check_item(checks.rows[0].cells[0], "Kontrolēts laidiens", "PR #73 apvienots pēc sekmīgas CI pārbaudes; publiski darbojas tieši apstiprinātais main komits 664b5e5.")
check_item(checks.rows[0].cells[0], "Pārbaudīta atkopšana", "Lietotne un tunelis atjaunojas pēc procesa kļūmes; praktiski izmēģināta atgriešanās uz iepriekšējo versiju.")
check_item(checks.rows[0].cells[1], "Minimāla publiskā virsma", "/api/health atklāj tikai stāvokli un laidiena SHA; publiskais notikumu API neizdod iekšējos laukus.")
check_item(checks.rows[0].cells[1], "Aizsargāta operatīvā piekļuve", "Centra diagnostika un datu atsvaidzināšana anonīmam lietotājam nav pieejama. MI nepublicē un nenosūta cilvēkus patstāvīgi.")

h = heading(doc, "Pašreizējā aizsardzības arhitektūra", size=14, before=5, after=3)
bottom_rule(h, color=LIGHT, size="6")
flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_p(flow, after=2, line=1.0)
parts = [
    ("Publiskais lietotājs", BLACK),
    ("  →  ", GRAY),
    ("Cloudflare", NAVY),
    ("  →  ", GRAY),
    ("drošs tunelis", NAVY),
    ("  →  ", GRAY),
    ("Atbalsts lietotne", NAVY),
    ("  →  ", GRAY),
    ("dati un MI pakalpojumi", BLACK),
]
for text, color in parts:
    r = flow.add_run(text)
    set_run(r, size=9.6, bold=color in (NAVY, BLACK), color=color)
flow_note = doc.add_paragraph()
flow_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_p(flow_note, after=5, line=1.05)
r = flow_note.add_run(
    "Cloudflare samazina tiešu servera ekspozīciju, taču pašreiz ir viena lokāla izcelsmes instance. "
    "Cloudflare Worker pašlaik neapkalpo publisko domēnu, un Vite priekšskatījuma serveris ir pieļaujams tikai demonstrācijai."
)
set_run(r, size=8.8, color=GRAY)

h = heading(doc, "Galvenie uzbrukuma ceļi", size=14, before=5, after=3)
bottom_rule(h, color=LIGHT, size="6")
attack_paths(doc)

h = heading(doc, "Ieviešanas robežas", size=14, before=5, after=3)
bottom_rule(h, color=LIGHT, size="6")
phases = doc.add_table(rows=1, cols=3)
phases.alignment = WD_TABLE_ALIGNMENT.CENTER
phases.autofit = False
hide_table_borders(phases)
for c in phases.rows[0].cells:
    set_cell_width(c, 58.5)
phase(phases.rows[0].cells[0], 1, "Demonstrācija", "Var izmantot tagad", GREEN, "Sintētiski dati, publiskā informācija, tikai lasāma karte un Centra skats. Kopīgā parole jānosauc kā apzināts demonstrācijas ierobežojums.")
phase(phases.rows[0].cells[1], 2, "Kontrolēts pilots", "Ar nosacījumiem", AMBER, "Pirms personas datu pilotprojekta vajadzīgi individuāli konti, MFA, novērtējums par ietekmi uz datu aizsardzību, glabāšanas noteikumi, failu karantīna un perimetra kontroles.")
phase(phases.rows[0].cells[2], 3, "Operatīva lietošana krīzē", "Vēl neatļaut", RED, "Nepieciešama augsta pieejamība, brīdinājumi, incidentu un rezerves kopiju mācības, kā arī obligātas misiju kompetences, riska, piekrišanas un četru acu pārbaudes.")

h = heading(doc, "Produktu nošķiršana", size=14, before=8, after=3)
bottom_rule(h, color=LIGHT, size="6")
product = doc.add_table(rows=1, cols=2)
product.alignment = WD_TABLE_ALIGNMENT.CENTER
product.autofit = False
hide_table_borders(product)
for c in product.rows[0].cells:
    set_cell_width(c, 88)
    set_cell_margins(c, top=20, start=0, bottom=0, end=140)

p = product.rows[0].cells[0].paragraphs[0]
style_p(p, after=2, line=1.04)
r = p.add_run("Sākt ar publisko informāciju un tikai lasāmu karti.  ")
set_run(r, size=9.2, bold=True, color=GREEN)
r = p.add_run("Tas ir drošākais atsevišķais produkts. Iedzīvotāju ziņojumus var pievienot kontrolētā pilotā pēc datu, failu un perimetra prasību izpildes.")
set_run(r, size=9.2)

p = product.rows[0].cells[1].paragraphs[0]
style_p(p, after=2, line=1.04)
r = p.add_run("Saglabāt izolētus vadības un nosūtīšanas moduļus.  ")
set_run(r, size=9.2, bold=True, color=RED)
r = p.add_run("Vadības centrs paliek demonstrācijas režīmā, bet misijas un brīvprātīgo nosūtīšanu nepalaiž līdz obligāto drošības pārbaužu ieviešanai. MI paliek tikai padomdevēja lomā.")
set_run(r, size=9.2)

doc.add_page_break()

# PAGE 2
title2 = doc.add_paragraph(style="Title")
remove_borders(title2._p)
style_p(title2, after=2, line=1.0, keep=True)
r = title2.add_run("Ieviestās prakses un atlikušais darbs")
set_run(r, size=20, bold=True)

intro2 = doc.add_paragraph()
style_p(intro2, after=6, line=1.1)
r = intro2.add_run(
    "Ekspluatācijas drošības gatavība ir 66%. Atlikušie 34 procentpunkti nav ievainojamību īpatsvars; "
    "tie apzīmē vēl neieviestās kontroles šajā iekšējā novērtējumā. Tās zemāk sasaistītas ar konkrētiem īpašniekiem un termiņiem."
)
set_run(r, size=9.4, color=GRAY)

practice = doc.add_table(rows=1, cols=2)
practice.alignment = WD_TABLE_ALIGNMENT.CENTER
practice.autofit = False
hide_table_borders(practice)
for c in practice.rows[0].cells:
    set_cell_width(c, 88)
    set_cell_margins(c, top=0, start=0, bottom=0, end=140)

left = practice.rows[0].cells[0]
right = practice.rows[0].cells[1]
left.paragraphs[0].clear()
right.paragraphs[0].clear()

p = left.paragraphs[0]
style_p(p, after=4, line=1.0)
r = p.add_run("JAU IEVIESTĀS LABĀS PRAKSES")
set_run(r, size=8.1, bold=True, color=GREEN)
p = right.paragraphs[0]
style_p(p, after=4, line=1.0)
r = p.add_run("PLĀNOTĀS KONTROLES ATLIKUŠAJIEM 34%")
set_run(r, size=8.1, bold=True, color=AMBER)

practice_item(left, "✓", "Kontrolēti laidieni un atkopšana", "PR apvienošana pēc CI, precīzs laidiena SHA, nemaināmas versijas, automātiska procesu atjaunošanās un praktiski pārbaudīta atgriešanās.", GREEN)
practice_item(left, "✓", "Minimāla publiskā datu virsma", "Veselības pārbaude izdod tikai statusu un SHA, atbildes netiek kešotas, bet publiskajā notikumu API izmantots atļauto lauku saraksts.", GREEN)
practice_item(left, "✓", "Piekļuves un ļaunprātīgas izmantošanas ierobežošana", "Aizsargātie maršruti anonīmam lietotājam dod 401; darbojas droši sīkfaili, izcelsmes pārbaudes un lietotnes pieprasījumu ierobežojumi.", GREEN)
practice_item(left, "✓", "Droša izstrāde, dati un faili", "CI izpilda 32 pārbaudes un kritisku atkarību kontroli. Servera atslēgas nav klientā; ievades validē, bet failus glabā privāti un pārbauda to īpašnieku, tipu un izmēru.", GREEN)
practice_item(left, "✓", "Cilvēka kontrole pār MI", "Ārkārtas informācija tiek apstrādāta pirms MI. Publicēšanu un cilvēku nosūtīšanu apstiprina cilvēks; MI atbildes ierobežo shēmas un avoti.", GREEN)

practice_item(right, "○", "Individuāla identitāte un atbildība", "Ieviest individuālus kontus, MFA, sesiju atsaukšanu, ierīču pārvaldību un iespēju noteikt, kurš operators veicis katru darbību.", AMBER)
practice_item(right, "○", "Perimetrs un repozitorija aizsardzība", "Ieviest Cloudflare WAF un ātruma ierobežojumus, pārskatīt Security Events un aizsargāt main ar obligātu PR, CI un CODEOWNERS apstiprinājumu.", AMBER)
practice_item(right, "○", "Datu un failu pilns dzīves cikls", "Pabeigt datu karti, NIDA, glabāšanas un dzēšanas noteikumus, minimālās datubāzes lomas, failu karantīnu un ļaunprogrammatūras skenēšanu.", AMBER)
practice_item(right, "○", "Noturīga ekspluatācija un incidentu vadība", "Aizstāt demonstrācijas serveri, pārbaudīt otru izcelsmes instanci, ieviest brīdinājumus, incidentu plānu un rezerves kopiju atjaunošanas mācības.", AMBER)
practice_item(right, "○", "Misiju un brīvprātīgo drošība", "Ieviest kompetences, riska un piekrišanas pārbaudes, četru acu principu un MI pretinieka scenāriju pārbaudes.", AMBER)

h = heading(doc, "Gatavība pa kontroles jomām", size=13, before=2, after=3)
scores = doc.add_table(rows=1, cols=2)
scores.alignment = WD_TABLE_ALIGNMENT.CENTER
scores.autofit = False
hide_table_borders(scores)
for c in scores.rows[0].cells:
    set_cell_width(c, 88)
    set_cell_margins(c, top=0, start=0, bottom=0, end=140)
left = scores.rows[0].cells[0]
right = scores.rows[0].cells[1]
left.paragraphs[0].clear()
right.paragraphs[0].clear()
compact_aspect(left, "Publiskā virsma", 92)
compact_aspect(left, "Izlaidumi un atkopšana", 90)
compact_aspect(left, "MI un ārkārtas situāciju loģika", 82)
compact_aspect(left, "Repozitorijs un piegādes ķēde", 74)
compact_aspect(left, "Dati un Supabase", 65)
compact_aspect(right, "Identitāte un sesijas", 58)
compact_aspect(right, "Faili", 55)
compact_aspect(right, "Uzraudzība un incidenti", 55)
compact_aspect(right, "Cloudflare perimetrs", 45)
compact_aspect(right, "Misijas un brīvprātīgie", 45)

h = heading(doc, "Četras nākamās prioritātes", size=14, before=2, after=3)
bottom_rule(h, color=LIGHT, size="6")
actions = doc.add_table(rows=2, cols=2)
actions.alignment = WD_TABLE_ALIGNMENT.CENTER
actions.autofit = False
hide_table_borders(actions)
for row in actions.rows:
    for c in row.cells:
        set_cell_width(c, 88)
        set_cell_margins(c, top=30, start=0, bottom=0, end=150)
action(actions.rows[0].cells[0], 1, "Perimetrs un koda pārvaldība", "Cloudflare un GitHub administratori", "17.09.2026", "Ieviest perimetra noteikumus un galvenā zara aizsardzību.", RED)
action(actions.rows[0].cells[1], 2, "Identitāte dati un faili", "Produkta privātuma un drošības īpašnieki", "Pirms personas datu pilota", "Ieviest kontus, MFA, NIDA, datu dzīves ciklu un failu karantīnu.", AMBER)
action(actions.rows[1].cells[0], 3, "Darbības nepārtrauktība", "Infrastruktūras īpašnieks", "10.10.2026", "Ieviest atbalstītu izpildvidi, otru instanci, brīdinājumus un atkopšanas mācības.", AMBER)
action(actions.rows[1].cells[1], 4, "Cilvēku un misiju drošība", "Produkta un operāciju vadība", "Pirms reālu misiju palaišanas", "Ieviest obligātās misiju drošības pārbaudes un saglabāt MI padomdevēja lomā.", RED)

h = heading(doc, "Atbilstības ceļš", size=14, before=3, after=3)
bottom_rule(h, color=LIGHT, size="6")
standards = doc.add_paragraph()
style_p(standards, before=5, after=1, line=1.09)
r = standards.add_run("Obligātā piemērojamība  ")
set_run(r, size=8.7, bold=True)
r = standards.add_run(
    "Personas datiem piemēro VDAR. Atsevišķi jānosaka Nacionālā kiberdrošības likuma, NIS2 un MK noteikumu Nr. 397 piemērojamība, "
    "Kibernoturības akta prasības katram tirgū laistam produktam ar digitāliem elementiem un MI akta prasības pēc MI faktiskās lomas."
)
set_run(r, size=8.7)

frameworks = doc.add_paragraph()
style_p(frameworks, after=2, line=1.08)
r = frameworks.add_run("Ieviešanas ietvars  ")
set_run(r, size=8.7, bold=True)
r = frameworks.add_run("NIST CSF 2.0 organizācijas risku pārvaldībai un OWASP ASVS 5.0 lietotnes kontrolēm. ISO/IEC 27001 sertifikāciju apsvērt, ja to prasa iepirkums vai tirgus.")
set_run(r, size=8.7)

sources = doc.add_paragraph()
style_p(sources, after=0, line=1.0)
r = sources.add_run("Primārie avoti  ")
set_run(r, size=7.4, bold=True, color=GRAY)
links = (
    ("VDAR", "https://eur-lex.europa.eu/eli/reg/2016/679/oj"),
    ("Nacionālās kiberdrošības likums", "https://likumi.lv/ta/id/353390-nacionalas-kiberdrosibas-likums"),
    ("MK Nr. 397", "https://likumi.lv/ta/id/361481-minimalas-kiberdrosibas-prasibas"),
    ("Kibernoturības akts", "https://digital-strategy.ec.europa.eu/en/policies/cra-summary"),
    ("MI akts", "https://eur-lex.europa.eu/eli/reg/2024/1689/oj"),
    ("NIST CSF 2.0", "https://www.nist.gov/cyberframework"),
    ("OWASP ASVS 5.0", "https://owasp.org/www-project-application-security-verification-standard/"),
)
for i, (label, url) in enumerate(links):
    if i:
        r = sources.add_run("  ·  ")
        set_run(r, size=7.4, color=GRAY)
    add_hyperlink(sources, label, url, size=7.4)

h = heading(doc, "Ieteicamais sapulces lēmums", size=14, before=8, after=3)
bottom_rule(h, color=LIGHT, size="6")
decision = doc.add_paragraph()
style_p(decision, after=3, line=1.11)
r = decision.add_run("Apstiprināt demonstrāciju pašreizējā vidē  ")
set_run(r, size=9.4, bold=True, color=GREEN)
r = decision.add_run("ar skaidri nosauktiem ierobežojumiem un sintētiskiem datiem. ")
set_run(r, size=9.4)
r = decision.add_run("Kontrolētu personas datu pilotu sākt tikai pēc 1. un 2. lēmuma izpildes. ")
set_run(r, size=9.4, bold=True, color=AMBER)
r = decision.add_run("Reālu krīzes operāciju un cilvēku nosūtīšanu atļaut tikai pēc 3. un 4. lēmuma izpildes un praktiskām incidentu mācībām.")
set_run(r, size=9.4, bold=True, color=RED)

doc.core_properties.title = "Atbalsts platformas drošības gatavība"
doc.core_properties.subject = "Vadības kopsavilkums krīzes centra vadībai"
doc.core_properties.author = "Sortium"
doc.core_properties.keywords = "Atbalsts kiberdrošība krīzes centrs Cloudflare VDAR NIS2 OWASP NIST"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
